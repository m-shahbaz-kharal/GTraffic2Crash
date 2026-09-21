"""Build the written report as a Word document.

Times New Roman throughout, 12 pt body, no italics, no header or footer, no
document metadata.  Bold marks headings and table column names; a single navy
is used to pick out the numbers that matter.
"""
import json
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIGS = os.path.join(ROOT, "docs", "doc")
OUT = os.path.join(ROOT, "docs", "doc", "GTraffic2Crash.docx")

FONT = "Times New Roman"
BLACK = RGBColor(0x00, 0x00, 0x00)
NAVY = RGBColor(0x1F, 0x38, 0x64)
GREY = RGBColor(0x59, 0x59, 0x59)

BANNED = "—–―−"      # em dash, en dash, horizontal bar, minus


# ------------------------------------------------------------------ helpers

def set_font(run, size=12, bold=False, color=BLACK):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = False
    run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = rpr.makeelement(qn("w:rFonts"), {})
        rpr.insert(0, rf)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(attr), FONT)


def check(text):
    for ch in BANNED:
        assert ch not in text, f"banned dash in: {text[:90]}"
    return text


def para(doc, pieces, size=12, space_after=7, align=None, style=None,
         space_before=0, keep_with_next=False):
    """pieces: a string, or a list of (text, bold, colour) tuples."""
    p = doc.add_paragraph(style=style)
    fmt = p.paragraph_format
    fmt.space_after = Pt(space_after)
    fmt.space_before = Pt(space_before)
    fmt.line_spacing = 1.10
    if align is not None:
        p.alignment = align
    if keep_with_next:
        fmt.keep_with_next = True
    if isinstance(pieces, str):
        pieces = [(pieces, False, BLACK)]
    for text, bold, colour in pieces:
        set_font(p.add_run(check(text)), size=size, bold=bold, color=colour)
    return p


def heading(doc, text, level=1):
    sizes = {0: 17, 1: 13.5, 2: 12}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0 if level == 0 else (16 if level == 1 else 12))
    p.paragraph_format.space_after = Pt(7 if level == 0 else 5)
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(check(text)), size=sizes[level], bold=True)
    return p


def bullets(doc, items, size=12):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.12
        pieces = item if isinstance(item, list) else [(item, False, BLACK)]
        for text, bold, colour in pieces:
            set_font(p.add_run(check(text)), size=size, bold=bold, color=colour)


def table(doc, rows, widths=None, size=11, highlight_col=None,
          highlight_row=None):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in enumerate(rows):
        for c, cell in enumerate(row):
            tc = t.cell(r, c)
            tc.text = ""
            p = tc.paragraphs[0]
            p.paragraph_format.space_before = Pt(2.5)
            p.paragraph_format.space_after = Pt(2.5)
            p.paragraph_format.line_spacing = 1.05
            colour = BLACK
            bold = (r == 0)
            if r > 0 and highlight_col is not None and c in highlight_col:
                colour = NAVY
                bold = True
            if highlight_row is not None and r == highlight_row:
                colour = NAVY
                bold = True
            set_font(p.add_run(check(str(cell))), size=size, bold=bold,
                     color=colour)
            if widths:
                tc.width = widths[c]
    if widths:
        for row in t.rows:
            for c, cell in enumerate(row.cells):
                cell.width = widths[c]
    return t


def figure(doc, name, width_in, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(os.path.join(FIGS, name), width=Inches(width_in))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(12)
    head, rest = caption.split("|", 1)
    set_font(cap.add_run(check(head)), size=10.5, bold=True, color=GREY)
    set_font(cap.add_run(check(rest)), size=10.5, color=GREY)


def spacer(doc, points=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(points)
    p.paragraph_format.space_before = Pt(0)
    set_font(p.add_run(""), size=4)


# -------------------------------------------------------------------- data

def load():
    ev = json.load(open(os.path.join(ROOT, "runs", "base_v1", "evaluation",
                                     "evaluation.json")))
    st = json.load(open(os.path.join(ROOT, "docs", "figures",
                                     "dataset_statistics.json")))
    return ev, st


def f3(x):
    return f"{x:.3f}"


# ------------------------------------------------------------------ writing

def build():
    ev, st = load()
    head = ev["headline"]
    bh = ev["by_horizon"]
    sc = ev["by_scenario"]
    counts = st["scenario_counts"]
    total_w = sum(counts.values())

    doc = Document()
    cp = doc.core_properties
    for prop in ("author", "last_modified_by", "title", "subject", "comments",
                 "category", "keywords"):
        try:
            setattr(cp, prop, "")
        except Exception:
            pass
    cp.revision = 1
    # The default python-docx template carries its own creation date; drop it
    # along with everything else, so the file says nothing about when or by whom
    # it was made.
    root = cp._element
    for name in ("created", "modified"):
        el = root.find("{http://purl.org/dc/terms/}" + name)
        if el is not None:
            root.remove(el)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    for s in doc.sections:
        s.top_margin = Inches(0.9)
        s.bottom_margin = Inches(0.9)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    # ------------------------------------------------------------- opening
    heading(doc, "Forecasting Town Traffic Five Seconds Ahead", level=0)
    para(doc,
         "A separate model already turns the pictures from 141 fixed traffic "
         "cameras into one top down view of a whole town, showing where every "
         "vehicle is. This work takes that view and asks the next question: "
         "given the last five seconds of it, what do the next five seconds look "
         "like? Crash warning, congestion detection and violation detection all "
         "sit on top of an answer to that, so the forecast is what gets built "
         "and measured here.")
    para(doc,
         "What follows runs in order: how the data was made, what the model is, "
         "how it was trained, and what it scores.")

    # ---------------------------------------------------------- 1. the data
    heading(doc, "1. Generating the data")
    para(doc,
         "Real roads cannot supply what this needs. Training requires the exact "
         "position of every vehicle at every instant, and it requires crashes "
         "often enough to learn from. Both are free in a simulator, so the data "
         "is generated in CARLA, an open driving simulator, using its built in "
         "town called Town05.")
    figure(doc, "fig_pipeline.png", 6.4,
           "Figure 1. |From simulator to training example. The town map is built "
           "once and reused by every episode.")

    heading(doc, "One episode", level=2)
    para(doc,
         "An episode is a single run of the simulator. Each one places between "
         "100 and 260 vehicles at random spawn points and lets them drive for "
         "450 seconds. Physics runs at 20 steps per second and the world is "
         "saved every tenth step, so one saved frame every 0.5 seconds. For "
         "every vehicle the record holds its position, heading, length and "
         "width, speed, class and behaviour tag, along with every collision the "
         "simulator reports. Episode number maps to configuration through a "
         "seeded generator, so any episode can be regenerated exactly. Half the "
         "episodes are ordinary traffic and half contain bad drivers.")

    heading(doc, "What nominal traffic means", level=2)
    para(doc,
         "Nominal does not mean identical. Every vehicle is driven by CARLA's "
         "Traffic Manager, and each one draws its own settings, so the traffic "
         "has natural spread rather than moving as a block:")
    bullets(doc, [
        "speed: drawn per vehicle between 5 percent above the posted limit for "
        "its lane and 20 percent below it",
        "following distance: drawn per vehicle between 2.0 and 4.5 metres",
        "traffic lights and road signs: always obeyed",
        "collision avoidance: fully on",
    ])
    para(doc,
         "Traffic like this still produces about one contact per simulated "
         "minute, almost all of it slow bumping in queues. That is the clean "
         "baseline the bad drivers are measured against.", space_before=4)

    heading(doc, "Adding bad drivers", level=2)
    para(doc,
         "On the other half of the episodes a fraction between 6 and 14 percent "
         "of the vehicles is drawn at random without replacement and reassigned "
         "to one of four profiles in rotation, so the four kinds appear in "
         "roughly equal numbers. Every profile works by changing the rules a "
         "driver follows rather than by making the driver blind, with one "
         "deliberate exception. A vehicle that cannot see other vehicles simply "
         "drives through them: an early calibration measured 126 collisions per "
         "simulated minute that way, which is carnage rather than traffic.")
    table(doc, [
        ["Profile", "The idea", "Exact settings"],
        ["Speeding",
         "Well over the limit with almost no gap, which is the classic rear end "
         "setup when the vehicle ahead stops for a light.",
         "speed 50 to 90 percent above the lane limit; following distance 0.3 to "
         "1.2 m; ignores stop and yield signs 90 percent of the time; obeys "
         "traffic lights; avoidance on"],
        ["Red light runner",
         "Enters junctions against the signal. Avoidance stays on, so a crash "
         "happens only when the cross traffic is genuinely unavoidable.",
         "ignores every traffic light and every sign; speed 20 to 50 percent "
         "above the limit; following distance 1.0 to 2.5 m; avoidance on"],
        ["Erratic, the impaired driver",
         "Weaves between lanes, judges gaps badly and is slow to notice other "
         "traffic. The only profile given partial blindness, because that is the "
         "behaviour being modelled rather than a side effect.",
         "random lane change 40 to 90 percent to the left and 40 to 90 percent "
         "to the right; ignores other vehicles 10 to 35 percent of the time; "
         "ignores lights 5 to 25 percent of the time; speed from 50 percent "
         "above the limit to 30 percent below it; following distance 0.2 to 1.5 m"],
        ["Tailgater",
         "Sits on the bumper of the vehicle ahead.",
         "following distance 0 m; speed 10 to 35 percent above the limit; obeys "
         "lights; avoidance on"],
    ], widths=[Inches(1.25), Inches(2.15), Inches(3.1)], size=10.5)
    para(doc, "Table 1. Every bad driver is a change of rules, not a change of "
              "physics.", size=10.5, space_before=4, space_after=10,
         align=WD_ALIGN_PARAGRAPH.LEFT)
    para(doc,
         "With this mix, an episode carrying bad drivers produces roughly one "
         "meaningful collision per simulated minute.")

    heading(doc, "Turning vehicles into a picture", level=2)
    para(doc,
         "The town is covered with square cells 4 metres across, 144 rows by 160 "
         "columns, which spans 576 by 640 metres and contains all of Town05. A "
         "cell does not store yes or no. It stores the fraction of its own area "
         "covered by vehicle bodies, a number from 0 to 1. Each vehicle "
         "rectangle is intersected with each cell exactly, by clipping polygons, "
         "so the value is a measured area rather than an estimate.")
    para(doc,
         "This is the most consequential choice in the whole dataset. Four "
         "metres is coarse next to a car, but because the stored value is an "
         "area fraction, a car straddling a boundary splits itself between the "
         "two cells in proportion. Its position inside the cell survives, and so "
         "does its motion from one frame to the next.")
    figure(doc, "fig_coverage.png", 5.9,
           "Figure 2. |The same car, two metres apart. Each cell is 4 metres "
           "across and holds the fraction of its area under a vehicle, so a "
           "shift far smaller than one cell still changes the numbers.")
    para(doc, "Three occupancy channels are kept, because the three groups move "
              "differently:")
    bullets(doc, [
        "cars",
        "large vehicles: trucks, vans and buses",
        "two wheelers: motorcycles and bicycles",
    ])
    para(doc,
         "Two more channels hold the speed of whatever is in the cell, east and "
         "north, in metres per second. These are targets only. The camera model "
         "that will feed this system produces occupancy and not speed, so "
         "handing the network true speeds as an input would flatter every number "
         "it later produces.", space_before=4)

    heading(doc, "The map the town sits on", level=2)
    para(doc,
         "Six further channels describe the town itself. They never change and "
         "are attached to every frame. The last two columns below are the actual "
         "six numbers stored for one open road cell and one cell inside a "
         "junction.")
    table(doc, [
        ["Channel", "Meaning", "Open road", "Junction"],
        ["road", "1 where a vehicle can be", "1.000", "1.000"],
        ["lane direction, north", "sine of the lane heading", "-1.000", "-1.000"],
        ["lane direction, east", "cosine of the lane heading", "-0.012", "0.022"],
        ["lane agreement", "1 where every lane in the cell points the same way",
         "1.000", "0.123"],
        ["junction", "1 inside an intersection", "0.000", "1.000"],
        ["speed limit", "posted limit in metres per second", "8.333", "8.333"],
    ], widths=[Inches(1.55), Inches(2.9), Inches(1.0), Inches(1.0)], size=10.5,
        highlight_col=(2, 3))
    para(doc, "Table 2. The six map channels, with a real sample of each.",
         size=10.5, space_before=4, space_after=10)
    para(doc,
         "Lane agreement is what separates the two samples. On a straight road "
         "every lane crossing the cell points the same way and the value is 1; "
         "inside a junction the lanes cross and it collapses to 0.123. About a "
         "quarter of the grid touches road at all, and telling the model where a "
         "vehicle can physically be saves it from spending capacity "
         "rediscovering the town.")
    para(doc,
         "Traffic light state is recorded but never given to the model. The task "
         "is to predict traffic from traffic. Handing over the signal phase "
         "would let the model learn that cars stop on red without learning that "
         "cars decelerate, and the camera model could not supply it anyway.")

    heading(doc, "Training examples and the split", level=2)
    para(doc,
         "A training example is a window of twenty consecutive frames: ten past "
         "frames the model sees, and ten future frames it has to produce. At 0.5 "
         "seconds a frame that is five seconds of history and five seconds of "
         "forecast. Windows slide forward one frame at a time, so neighbouring "
         "windows overlap heavily. That overlap is why the train and validation "
         "split is made by episode and never by window: 53 of the 355 episodes "
         "are held out whole, and no frame from a training episode can appear in "
         "validation.")
    pct = {k: 100.0 * v / total_w for k, v in counts.items()}
    table(doc, [
        ["Quantity", "Value"],
        ["Episodes", "355, of which 302 are used for training and 53 held out"],
        ["Frames", f"{st['frames']:,} at 0.5 s, which is "
                   f"{st['frames'] * st['frame_dt_s'] / 3600:.0f} hours of traffic"],
        ["Vehicles on screen", f"{st['vehicles_per_frame_mean']:.0f} per frame on average"],
        ["Grid", "144 by 160 cells of 4 m, so 576 by 640 metres"],
        ["Channels per frame", "3 vehicle classes, 2 speed channels, 6 map channels"],
        ["Training windows", "264,742"],
        ["Validation windows", "9,381, taken every fifth frame so they are not "
                               "near duplicates of each other"],
        ["Occupancy", f"{100 * st['occupied_cell_fraction_mean']:.1f} percent of "
                      "cells hold a vehicle at any moment"],
        ["Window labels",
         f"nominal {pct['nominal']:.0f} percent, incident traffic "
         f"{pct['incident_traffic']:.0f} percent, collision {pct['collision']:.0f} "
         f"percent, low speed contact {pct['low_speed_contact']:.0f} percent"],
    ], widths=[Inches(1.8), Inches(4.65)], size=10.5)
    para(doc, "Table 3. The dataset at a glance.", size=10.5, space_before=4,
         space_after=10)
    para(doc,
         "The window labels are for reporting only and never reach the model. A "
         "window counts as a collision if a crash at 3 metres per second or more "
         "happens inside its ten future frames, as low speed contact if a slower "
         "contact happens, as incident traffic if bad drivers are present but "
         "nothing is hit, and as nominal otherwise.")

    # --------------------------------------------------------- 2. the model
    heading(doc, "2. The model")
    heading(doc, "Where the design comes from", level=2)
    para(doc,
         "Two papers settled the shape of it.")
    para(doc,
         "OpenSTL (arXiv 2306.11249) benchmarks 24 spatio temporal prediction "
         "models on a common footing. On TaxiBJ, a city traffic grid and the "
         "closest public relative of this problem, the best model without "
         "recurrence matches the best recurrent one while using about one "
         "twenty sixth of the computation. Their reading is that recurrence pays "
         "off on fast changing data and wastes effort on slow changing data. "
         "Town traffic at two frames per second is slow changing.")
    para(doc,
         "Earthformer (arXiv 2207.05833) compares seven ways of factorising "
         "attention across space and time under identical conditions. Attention "
         "applied along one axis at a time, plus a small set of global vectors, "
         "wins on every measure while being among the cheapest of the seven. "
         "Attending over the whole space and time block at once is the worst.")
    para(doc,
         "So the model is an encoder and a decoder built from axial attention "
         "blocks with global vectors, and it produces all ten future frames in "
         "one pass instead of one frame at a time, which removes any chance of "
         "error compounding down the horizon.")

    heading(doc, "What we changed", level=2)
    bullets(doc, [
        [("Soft targets. ", True, BLACK),
         ("Earthformer predicts real valued weather fields. Here the target is "
          "an area fraction bounded between 0 and 1, which is what makes the "
          "loss in section 3 well posed.", False, BLACK)],
        [("A convolutional stem and head instead of patch embedding. ", True, BLACK),
         ("A 16 metre token cannot say which 4 metre cell a car is in, so the "
          "stem keeps full resolution features and hands the last observed "
          "frame's features straight across to the head.", False, BLACK)],
        [("Start from persistence. ", True, BLACK),
         ("The head adds its output to the last observed frame rather than "
          "drawing the future from nothing. At step one the model is exactly "
          "“nothing moves”, so every parameter it spends goes on the "
          "change.", False, BLACK)],
        [("A second output for speed, ", True, BLACK),
         ("and a loss term that ties the two outputs together.", False, BLACK)],
        [("Map conditioning. ", True, BLACK),
         ("Weather data has no road network. Town traffic does, and the road "
          "network is most of the prior.", False, BLACK)],
    ])

    heading(doc, "The architecture", level=2)
    figure(doc, "fig_training.png", 6.5,
           "Figure 3. |One training step. The upper row is the network, the "
           "lower row is the objective it is scored against.")
    para(doc,
         "The stem turns each 144 by 160 frame into a 36 by 40 grid of tokens of "
         "width 640, so the encoder is working on 14,400 tokens across the ten "
         "input frames. Attending over all of them at once would be 207 million "
         "pairs in every layer. Attending along one axis at a time means each "
         "token reaches 86 others instead of 14,400, and the eight global "
         "vectors put back the long range links the factorisation gives up.")
    table(doc, [
        ["Part", "Parameters"],
        ["Convolution stem", "2.0 M"],
        ["Axial encoder, 12 blocks", "236.3 M"],
        ["Axial decoder, 8 blocks", "197.0 M"],
        ["Convolution head", "3.3 M"],
        ["Position and global embeddings", "0.9 M"],
        ["Total", "439.5 M"],
    ], widths=[Inches(2.6), Inches(1.5)], size=10.5)
    para(doc, "Table 4. Where the capacity sits.", size=10.5, space_before=4,
         space_after=10)

    # ------------------------------------------------------- 3. training it
    heading(doc, "3. Training")
    heading(doc, "Batches", level=2)
    para(doc,
         "A batch is 8 windows on each of two H100 GPUs, so 16 windows per step, "
         "drawn at random from the training episodes. Mirroring is the only "
         "augmentation: with probability one half the whole batch is mirrored "
         "east to west, and independently with probability one half north to "
         "south. The map and the validity mask are mirrored with it and the sign "
         "of the matching speed channel is flipped, so the world stays "
         "consistent. It is applied to the whole batch at once because the map "
         "is shared across the batch, and mirroring one window's traffic without "
         "its map would put cars on the wrong side of the road.")

    heading(doc, "What goes in and what comes out", level=2)
    para(doc,
         "Each of the ten input frames carries nine channels: three occupancy "
         "channels and the six map channels. A single cell of a single input "
         "frame is therefore nine numbers. For a cell holding most of one car on "
         "an open road they are:")
    table(doc, [
        ["cars", "large", "two wheelers", "road", "lane north", "lane east",
         "lane agreement", "junction", "speed limit"],
        ["0.386", "0.000", "0.000", "1.000", "-1.000", "-0.012", "1.000",
         "0.000", "0.417"],
    ], widths=[Inches(0.56), Inches(0.56), Inches(0.72), Inches(0.52),
               Inches(0.64), Inches(0.60), Inches(0.90), Inches(0.66),
               Inches(0.64)], size=9.5, highlight_col=(0,))
    para(doc, "Table 5. One cell of one input frame. The posted limit is divided "
              "by 20 before it is fed in, so 8.333 metres per second arrives as "
              "0.417 and sits in the same range as everything else.",
         size=10.5, space_before=4, space_after=10)
    para(doc,
         "The whole input is ten of those frames stacked, so 10 by 9 by 144 by "
         "160 numbers. The output is 10 by 3 by 144 by 160 occupancy values and "
         "10 by 2 by 144 by 160 speed values: for every future frame, the "
         "fraction of each cell covered by each of the three vehicle groups, and "
         "how fast whatever is in that cell is going.")

    heading(doc, "The loss and why each part is there", level=2)
    para(doc,
         "The objective has four parts, shown along the bottom of Figure 3.")
    para(doc, [("Occupancy loss. ", True, BLACK),
                ("Cell by cell cross entropy between the predicted fraction and "
                 "the true fraction, counting cells at or beside a vehicle ten "
                 "times a cell of open ground. The target is a fraction, so the "
                 "loss has to reach its lowest point exactly at the true "
                 "fraction. Cross entropy does. The two obvious alternatives do "
                 "not:", False, BLACK)])
    bullets(doc, [
        "Dice and intersection over union: for a cell that is half covered, soft "
        "Dice is lowest when the prediction is 1.0, so they push every partly "
        "covered cell to full and inflate the vehicle count (arXiv 2303.16296).",
        "Focal loss: not a proper scoring rule for any positive focusing "
        "parameter, and its bias runs toward under confidence (arXiv 2011.09172), "
        "which here means quietly losing vehicles.",
    ])
    para(doc,
         "The usual worry about cross entropy is that 98 percent of cells are "
         "empty and will drown the rest. That is true at the very first step, "
         "where the network says one half everywhere and 23,000 empty cells "
         "outvote 460 occupied ones, and it is fixed where it happens: the "
         "output bias starts at the measured occupancy rate (arXiv 1708.02002), "
         "instead of the loss being bent around the problem.", space_before=4)
    para(doc, [("Speed loss. ", True, BLACK),
                ("Absolute error between the predicted and true speed, on "
                 "occupied cells only. An empty cell has no speed to be right "
                 "about.", False, BLACK)])
    para(doc, [("Movement loss. ", True, BLACK),
                ("Take the occupancy that was actually observed, slide it along "
                 "the speeds the model predicted, and require the result to "
                 "match the next frame. This is the term that carries the "
                 "physics: traffic has to move from where it was, it cannot "
                 "appear in one place and evaporate from another (arXiv "
                 "2203.03312).", False, BLACK)])
    para(doc, [("Count loss. ", True, BLACK),
                ("The total predicted vehicle area in each frame, for each "
                 "class, has to match the truth. Counting cells near vehicles "
                 "ten times higher is useful, but it stops the objective being "
                 "strictly proper, so the total will drift. This term measures "
                 "that drift in the one quantity that is physically meaningful, "
                 "the number of vehicles, and pulls it back (arXiv 2009.13077).",
                 False, BLACK)])
    para(doc,
         "The four are combined with weights 1.0, 0.5, 0.5 and 0.001. The count "
         "weight looks tiny and is not: the first three are averages over "
         "roughly 1.7 million weighted cells, while the count is a sum over the "
         "same cells, so its natural gradient at any one cell is about a "
         "thousand times larger. Matching the two at an occupied cell is what "
         "puts the weight near 0.001.")

    heading(doc, "Settings", level=2)
    table(doc, [
        ["Setting", "Value"],
        ["Optimiser", "AdamW, betas 0.9 and 0.95, weight decay 0.05 on weight "
                      "matrices only"],
        ["Learning rate", "2e-4, with 1,000 warm up steps and then a cosine "
                          "decay to 5 percent"],
        ["Precision", "bfloat16 mixed, with activation checkpointing on every "
                      "block"],
        ["Batch", "8 windows per GPU on 2 GPUs, so 16 per step"],
        ["Length", "27,297 steps of a 32,000 step budget, which is 1.65 passes "
                   "over the training set and 35 hours of wall clock"],
    ], widths=[Inches(1.5), Inches(4.95)], size=10.5)
    para(doc, "Table 6. Training settings.", size=10.5, space_before=4,
         space_after=10)

    # ------------------------------------------------------------ 4. result
    heading(doc, "4. Results")
    para(doc,
         "Everything below is measured on the 53 held out episodes, 9,381 "
         "windows, none of which were seen during training.")
    heading(doc, "How it is scored", level=2)
    para(doc,
         "The headline number is soft intersection over union. For every valid "
         "cell of every future frame, add up the smaller of the predicted and "
         "the true coverage, add up the larger, and divide one by the other. A "
         "perfect forecast scores 1.0. It is computed on the fractions directly "
         "with no thresholding, and the cells outside the town boundary, which "
         "are 29 percent of the grid and always empty, are excluded so that "
         "blank ground cannot flatter the score.")
    heading(doc, "The three baselines", level=2)
    bullets(doc, [
        [("Repeat the last frame. ", True, BLACK),
         ("Copy the present into the future. It sounds trivial, but about half "
          "the town is stopped at any moment because of the signals, so it is "
          "right about half the time.", False, BLACK)],
        [("Constant speed, true speeds given. ", True, BLACK),
         ("Slide every vehicle along its real velocity, held fixed across the "
          "horizon. The model never sees those velocities, so this baseline is "
          "handed information the model does not have. It answers how far pure "
          "extrapolation gets you even with perfect knowledge of current "
          "motion.", False, BLACK)],
        [("Constant speed, speeds estimated. ", True, BLACK),
         ("The same thing, but the velocities are estimated from the last two "
          "grids with off the shelf optical flow. This is the version that could "
          "actually be deployed.", False, BLACK)],
    ])
    para(doc,
         "All three use the same mass conserving warp as the training objective, "
         "so a difference in score is a difference in prediction and not in "
         "resampling.", space_before=4)
    table(doc, [
        ["", "Over the five seconds", "At five seconds"],
        ["Our model", f3(head["model"]["soft_iou_mean"]),
         f3(head["model"]["soft_iou_last_step"])],
        ["Constant speed, true speeds given",
         f3(head["constant_velocity_oracle_flow"]["soft_iou_mean"]),
         f3(head["constant_velocity_oracle_flow"]["soft_iou_last_step"])],
        ["Repeat the last frame",
         f3(head["persistence"]["soft_iou_mean"]),
         f3(head["persistence"]["soft_iou_last_step"])],
        ["Constant speed, speeds estimated",
         f3(head["constant_velocity_estimated_flow"]["soft_iou_mean"]),
         f3(head["constant_velocity_estimated_flow"]["soft_iou_last_step"])],
    ], widths=[Inches(2.8), Inches(1.85), Inches(1.5)], size=10.5,
        highlight_row=1)
    para(doc, "Table 7. Agreement with what actually happened, higher is "
              "better.", size=10.5, space_before=4, space_after=10)
    figure(doc, "fig_horizon.png", 5.9,
           "Figure 4. |The gap widens with the horizon, which is where a "
           "forecast is worth having.")
    para(doc,
         f"At half a second there is little to win: copying the present already "
         f"scores {f3(bh['persistence']['soft_iou'][0])} and the extrapolation "
         f"given true speeds scores "
         f"{f3(bh['constant_velocity_oracle_flow']['soft_iou'][0])}. By two "
         f"seconds that extrapolation has fallen to "
         f"{f3(bh['constant_velocity_oracle_flow']['soft_iou'][3])} because "
         f"constant velocity cannot turn a corner, stop at a light or join the "
         f"back of a queue, while the model is still at "
         f"{f3(bh['model']['soft_iou'][3])}. The skill the model has and the "
         f"baselines do not is telling the vehicle that will sit still for five "
         f"seconds apart from the one that will travel thirty six metres, and "
         f"that cannot be read off a single frame.")

    heading(doc, "The score is not bought by hedging", level=2)
    para(doc,
         "There is a known failure in this literature where a model wins overlap "
         "scores by predicting less area than really exists, and so is never "
         "wrong in an expensive way. The check is the ratio of predicted total "
         "vehicle area to true total vehicle area, which should be 1. Ours runs "
         f"from {f3(bh['model']['bias'][0])} half a second out to "
         f"{f3(bh['model']['bias'][-1])} at five seconds, and the predicted "
         f"vehicle count stays within "
         f"{100 * head['model']['vehicle_count_relative_error_mean']:.1f} percent "
         "of the truth. The forecast is very slightly generous rather than "
         "timid, which is the right side to err on.")

    heading(doc, "Where it holds up", level=2)
    table(doc, [
        ["Window type", "Agreement"],
        ["Nominal traffic", f3(sc["nominal"]["model"]["soft_iou_mean"])],
        ["Incident traffic, bad drivers present",
         f3(sc["incident_traffic"]["model"]["soft_iou_mean"])],
        ["Low speed contact", f3(sc["low_speed_contact"]["model"]["soft_iou_mean"])],
        ["A crash happens in the forecast window",
         f3(sc["collision"]["model"]["soft_iou_mean"])],
    ], widths=[Inches(3.6), Inches(1.4)], size=10.5, highlight_col=(1,))
    para(doc, "Table 8. Accuracy by kind of window.", size=10.5, space_before=4,
         space_after=10)
    gap = (sc["nominal"]["model"]["soft_iou_mean"]
           - sc["collision"]["model"]["soft_iou_mean"]) * 100
    para(doc,
         f"The hardest windows, the ones that contain an actual crash, cost "
         f"about {gap:.1f} points against ordinary traffic. Accuracy bends where "
         "it matters rather than breaking.")

    heading(doc, "What the forecast looks like", level=2)
    figure(doc, "fig_forecast.png", 6.2,
           "Figure 5. |Five seconds ahead on a held out episode. Dark blue is "
           "cars, rust is trucks and buses, olive is two wheelers, and grey is "
           "the drivable road.")
    para(doc,
         "The queues at the junctions are in the right places and roughly the "
         "right lengths, worked out from the behaviour of the traffic itself "
         "since the model is never told the signal phase. Vehicles stay on the "
         "road. Where the forecast is fainter than the truth it is genuinely "
         "uncertain, usually about whether a queue has started moving yet, and "
         "the total mass still matches.")

    # ----------------------------------------------------------- 5. what is next
    heading(doc, "5. What comes next")
    bullets(doc, [
        [("Finish the run. ", True, BLACK),
         ("The numbers above come from 1.65 passes over the training data and "
          "the curve has not flattened.", False, BLACK)],
        [("Crash anticipation. ", True, BLACK),
         ("Read the forecast at the cell where a collision is about to happen "
          "and compare it against a matched control site, by lead time. That "
          "turns the forecast into a measured number of seconds of warning, "
          "which is the number a deployment would actually be judged on.",
          False, BLACK)],
        [("A second town. ", True, BLACK),
         ("Everything so far is Town05, so some of the score is certainly "
          "memorised layout. Measuring how much means training on one town and "
          "testing on another.", False, BLACK)],
        [("Camera input end to end. ", True, BLACK),
         ("Drive the forecaster from the camera model's grid rather than the "
          "simulator's. The forecaster was built to take exactly that format and "
          "was never given anything the camera model cannot produce, so the swap "
          "should be mechanical; what the input noise costs is the open "
          "question.", False, BLACK)],
        [("Longer horizon and finer cells. ", True, BLACK),
         ("Five seconds at 4 metres is the first version. Both are "
          "preprocessing constants rather than architecture, and the raw vehicle "
          "states are kept, so either can be revised for the cost of one pass "
          "over the data.", False, BLACK)],
    ])

    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    build()
